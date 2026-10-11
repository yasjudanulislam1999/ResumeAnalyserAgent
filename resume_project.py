from pathlib import Path
import os
from dotenv import load_dotenv
from groq import Groq
from typing import Optional
from pydantic import BaseModel
import json
from pypdf import PdfReader
from docx import Document
import time


load_dotenv()

class Jobdescription(BaseModel):
    role: str
    required_skills: str
    preferred_skills : str
    minimum_experience: float | None
    education_requirements : list[str]
    responsibilities: list[str]


Jobdescription_schema = Jobdescription.model_json_schema()

system_prompt = f'''
You are an expert HR assistant your job is to analyze the job description and extract structured information from it

Return ONLY valid JSON matching this schema:

{Jobdescription_schema}

IMPORTANT:
DO NOT RETURN THE SCHEMA ITSELF
DO NOT RETURN FIELDS LIKE "Properties", "title", "type".
Fill the schema with actual information extracted from the job description

If minimum information is not mentioned, return null
If an information for a list is missing, return an empty list.
DO NOT invent any information
'''

job_description = '''About The Role

We are seeking a highly skilled and motivated Data Analyst to join our dynamic Underwriting Technical Transformation team within the Health Insurance Line of business. This role offers a unique opportunity to be at the forefront of driving innovation and data-driven decision-making in underwriting and pricing.

As a Data Analyst, you will play a pivotal role in shaping our centralized data analytics function, which supports technical areas such as underwriting, pricing, and commercial. You will work collaboratively with cross-functional teams, including Data Science, Data Program, Portfolio Monitoring, Pricing and Underwriting, to design, build and implement robust data pipelines in order to deliver actionable insights that enhance business performance and customer outcomes.

What You Will Do

Support and participate in the build and delivery of a clear roadmap for the creation and modernization of a robust data environment aimed at ensuring commercial and strategic success for the business
Technically support the design, construction and maintenance of data for the business
Understand and implement database and data engineering concepts and terminology, design principles and elements in order to deliver data and analytics
Build and maintain production-grade data pipelines in order to automate data extraction, transformation, and loading (ETL) processes
Coordination of data projects with different levels of complexity depending on existing methods, processes and tools
Collaborate and engage with Data Program and Data Science teams to support the creation, maintenance and enhancement of reports, tools, datasets and dashboards to facilitate analysis and ongoing review of KPIs, e.g. leveraging Power BI
Understand the data relationships that exist between the business processes and systems. In turn apply this knowledge when interpreting the information requirements from the business
Perform in-depth data analysis to identify trends, patterns, and opportunities for optimization in pricing & underwriting processes.
Deliver data outputs into a format for presentation to senior management
Gather, interpret, identify gaps and translate business MI and reporting requirements into solutions
Manages and coordinates stakeholders and builds a robust communication structure for data-related projects
Advises stakeholders and actively clarifies their expectations and requirements
Ability to assume responsibility for change, communication, quality and risk management
AI Readiness Work effectively in an environment shaped by artificial intelligence (AI), machine learning, data, analytics and cloud-based tools, using insights responsibly with our standards of data governance, security and ethical use.

What You Will Need

Bachelor or above in a quantitative discipline (e.g. Mathematics, Computer Science) from top tier university
Proven relevant experience (min. 3 years) in data analysis, preferably within the insurance or healthcare sector.
Hands-on experience in writing maintainable and testable production-grade Python in a collaborative approach
Strong proficiency with direct exposure to SQL/ SQL optimization, Git version control, CI/CD practices as well as data visualization tools (e.g. Power BI). VBA would be useful
Excellent analytical skills (capable of understanding complex data structures, organize and structure data extracts) with the ability to translate complex data into actionable insights.
Works with little guidance and direction and performs majority of tasks independently.
High attention to detail including a commitment to accuracy of work.
Excellent communication and collaboration skills, with the ability to explain complicated processes and concepts to non-experts and to work effectively with diverse cross-functional stakeholders.
A proactive mindset with a passion for driving change and innovation through data.

Why join us

This is more than just a data analytics role – it’s an opportunity to make a tangible impact on the future of health insurance underwriting. You’ll be part of a team that values creativity, collaboration, and continuous improvement, all while working in a fast-paced and evolving industry.

If you’re ready to take on a challenging and rewarding role where your expertise in data analytics can truly make a difference, we’d love to hear from you!
'''

user_prompt = f'''Analyze the job description and : {job_description}'''

message_system = {
    "role": 'system',
    "content": system_prompt
}

message_user = {
    'role': 'user',
    "content": user_prompt
}



messages = [message_system,message_user]

my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("Groq API Key not found")

client = Groq(api_key=my_api_key)

model = "openai/gpt-oss-120b"

response_format = {
    "type":"json_object"
}

response = client.chat.completions.create(model=model,messages=messages,response_format=response_format)

job = Jobdescription.model_validate_json(response.choices[0].message.content)

#part resume

class MatchResult(BaseModel):
    score : float
    details : dict

class Experience(BaseModel):
    company : str| None=None
    role : str | None=None
    duration : str | None=None
    description : str | None=None
    skills_used : list[str] = []

class Resume(BaseModel):
    name : str | None=None
    role : str | None = None
    duration : str | None = None
    total_experience_years : float | None=None

    skills : list[str] = []
    experiences : list[Experience] = []
    education : list[str]=[]
    projects : list[str] = []
    certifications : list[str] = []

resume_schema = Resume.model_json_schema()

def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    message={
        "role": "user",
        "content" : prompt
    }
    messages=[message]
    response_format={
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)

def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume

from pypdf import PdfReader
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None



# lets do it now
resume_folder = Path("resumes")
all_results=[]
for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume=parse_resume(resume_text) # llm call1
    time.sleep(5)
    result = final_score(job, parsed_resume) #llm caLL2
    #score and details
    #acount chtgpt
    # request bhejna shhur krega millions
    #chattgot server jam ho jayega
    time.sleep(5)
    print("Score:", result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })
all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)
top_2 = all_results[:2]
worst_2 = all_results[-2:]


print("TOP 2 CANDIDATES")
for candidate in top_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])

print("LOWEST 2 CANDIDATES")
for candidate in worst_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])







